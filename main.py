"""
Download de notas fiscais em lote — SAP + portal web + verificação.

Lê a planilha de ocorrências do mês e baixa o PDF de cada nota, com menu no
terminal:

    1 - Baixar notas de SERVIÇO no SAP (transação NIXNFSE)
    2 - Baixar notas de MATERIAL no portal web (a partir do relatório do SAP)
    3 - Verificar planilha × pasta
    4 - Fluxo completo (1 → 2 → 3)
    5 - Rebaixar só o que está faltando (3 → 1 → 2 → 3)
    6 - Trocar a planilha de entrada (só nesta sessão)

Todas as etapas leem a planilha pela MESMA função (ler_planilha), então o
nome do arquivo de cada linha é sempre o mesmo em todo o fluxo:
  1ª vez que a nota aparece → '12345.pdf'
  2ª vez                    → '12345 (2).pdf'  ... e assim por diante.

Configuração (planilha, pasta de destino e URL do portal) no arquivo .env —
veja o .env.example.
"""
import os
import re
import time
import shutil
import tempfile
import traceback
import unicodedata
from collections import defaultdict
from datetime import datetime, timedelta

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string


# =====================================================================
#                       CONFIGURAÇÕES GERAIS
# =====================================================================
try:                                   # lê o .env, se o python-dotenv estiver instalado
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

# Planilha de entrada (a que VOCÊ atualiza todo mês). Dá para trocar em
# tempo de execução pela opção 6 do menu, sem editar o código.
ARQUIVO_LISTA = os.getenv("ARQUIVO_LISTA", "").strip().strip('"')
PASTA_DESTINO = os.getenv("PASTA_DESTINO", os.path.join(os.getcwd(), "notas")).strip().strip('"')

ABA = None           # None = aba ativa

# ---------- COLUNAS DA PLANILHA ----------
# O script procura as colunas PELO NOME DO CABEÇALHO (em qualquer posição e
# em qualquer uma das primeiras LINHAS_BUSCA_CABECALHO linhas), mostra o que
# achou e pede confirmação. A planilha original NÃO é alterada.
# Se o cabeçalho da sua planilha tiver outro nome, acrescente aqui (sem se
# preocupar com acento ou maiúscula).
NOMES_COL_NOTA = ["nota fiscal", "nota", "nf", "nfs-e", "nfse", "numero nf",
                  "nº nf", "numero da nota", "n nota", "nº nota"]
NOMES_COL_OCOR = ["ocorrência", "ocorrencia", "wf gaia", "ocorrência gaia",
                  "id ocorrência", "nº ocorrência"]
# Cabeçalhos com estas palavras nunca são escolhidos (ex.: "Valor da Nota").
PALAVRAS_IGNORADAS = ["valor", "data", "serie", "chave", "status", "tipo",
                      "vencimento", "emissao", "cnpj"]
LINHAS_BUSCA_CABECALHO = 10

# Usados só se o cabeçalho NÃO for encontrado pelo nome.
COL_NOTA = "D"
COL_OCOR = "H"
LINHA_INICIO = 2

# Processa só as N primeiras notas na etapa do SAP (teste). 0 = todas.
LIMITE_NOTAS = 0

# Arquivos gerados na pasta de destino
RELATORIO_EXCEL     = os.path.join(PASTA_DESTINO, "_relatorio_nfse.xlsx")
ARQUIVO_VERIFICACAO = os.path.join(PASTA_DESTINO, "_verificacao_nfse.xlsx")
PASTA_DIAGNOSTICO   = os.path.join(PASTA_DESTINO, "_diagnostico_nix")


# =====================================================================
#                       CONFIGURAÇÕES — SAP
# =====================================================================
# ---------- MODO SEGUNDO PLANO (captura do arquivo temporário) ----------
# Deixando None, o script descobre sozinho as pastas temporárias, mede o
# tempo de varredura e descarta as lentas. Depois da 1ª nota capturada,
# trava na pasta que funcionou. Para fixar:
#   PASTAS_MONITORADAS = [r"C:\Users\SEU_USUARIO\AppData\Local\Temp"]
PASTAS_MONITORADAS = None
PROFUNDIDADE_VARREDURA = 2
LIMITE_TEMPO_VARREDURA = 1.5

TIMEOUT_CAPTURA   = 8.0
INTERVALO_CAPTURA = 0.2
TAMANHO_MAXIMO_PDF = 50 * 1024 * 1024
TAMANHO_MINIMO_PDF = 2048
LEITURAS_ESTAVEIS  = 2

# Se a captura falhar, tenta o método antigo (3 pontos + Salvar como).
# Quando dispara, precisa da tela livre naquele instante.
USAR_FALLBACK_PYAUTOGUI = False

TRES_PONTOS = (800, 160)
SETAS_ATE_SALVAR = 2
ESPERA_RENDER = 3.0
ESPERA_MENU   = 0.8
TIMEOUT_PDF    = 8.0
INTERVALO_POLL = 0.2

# (EXPERIMENTAL) reaproveita a tela de seleção do NIXNFSE entre as notas
REUSAR_TELA_SELECAO = False
VOLTAS_ATE_SELECAO = 2

PARCIAL_A_CADA = 10
TIPOS_PDF_ACEITOS = ["PDF Files", "PDF (*.pdf)", "Arquivos PDF", "PDF", "Todos os arquivos (*.*)"]


# =====================================================================
#                       CONFIGURAÇÕES — NIX WEB
# =====================================================================
# Página de busca de ocorrências do portal web (fica no .env)
URL_NIX  = os.getenv("NIX_URL", "")
PERFIL_NAVEGADOR = os.path.join(os.environ.get("USERPROFILE", os.path.expanduser("~")),
                                "AppData", "Local", "nix_playwright_profile")
ESPERA_BUSCA = 10.0
TIMEOUT_RESPOSTA_PDF = 15000  # ms
COLUNA_OCORRENCIA = "WF GAIA"
# =====================================================================


# =====================================================================
#            IMPORTS "SOB DEMANDA" (só quando a etapa é usada)
# =====================================================================
# Assim a verificação roda mesmo sem SAP aberto, e o menu abre na hora
# (pywinauto e playwright demoram para importar).
win32com = pyautogui = pyperclip = Desktop = None
sync_playwright = PlaywrightTimeout = None


def _carregar_libs_sap():
    global win32com, pyautogui, pyperclip, Desktop
    if Desktop is not None:
        return
    import win32com.client          # noqa: F401  (vincula o nome global win32com)
    import pyautogui as _pyautogui
    import pyperclip as _pyperclip
    from pywinauto import Desktop as _Desktop
    pyautogui, pyperclip, Desktop = _pyautogui, _pyperclip, _Desktop
    pyautogui.PAUSE = 0.02


def _carregar_libs_nix():
    global sync_playwright, PlaywrightTimeout
    if sync_playwright is not None:
        return
    from playwright.sync_api import sync_playwright as _sp, TimeoutError as _te
    sync_playwright, PlaywrightTimeout = _sp, _te


# =====================================================================
#                       FUNÇÕES COMUNS
# =====================================================================
FILL_VERDE    = PatternFill("solid", fgColor="C6EFCE")
FILL_AMARELO  = PatternFill("solid", fgColor="FFEB9C")
FILL_VERMELHO = PatternFill("solid", fgColor="FFC7CE")
FILL_TITULO   = PatternFill("solid", fgColor="1F3864")
FILL_CABEC    = PatternFill("solid", fgColor="2E75B6")
FONTE_TITULO  = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
FONTE_CABEC   = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
FONTE_NORMAL  = Font(name="Calibri", size=11)
BORDA = Border(left=Side(style="thin"), right=Side(style="thin"),
               top=Side(style="thin"), bottom=Side(style="thin"))
CENTRO = Alignment(horizontal="center", vertical="center")


def limpar(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def tem_digitos(valor):
    """Separa dado real de cabeçalho/legenda ('Ocorrência', 'NF'...)."""
    return bool(re.search(r"\d", limpar(valor)))


def extrair_ocorrencia(valor):
    """'4800000001 - FIN0000001' → '4800000001' (o SAP não aceita o formato composto)."""
    valor = limpar(valor)
    m = re.match(r"^(\d+)\s*-\s*FIN\d*", valor, re.IGNORECASE)
    if m:
        return m.group(1)
    return valor


def pdf_valido(caminho):
    """True se o arquivo começa com a assinatura %PDF-."""
    try:
        if not os.path.exists(caminho) or os.path.getsize(caminho) == 0:
            return False
        with open(caminho, "rb") as f:
            return f.read(5) == b"%PDF-"
    except Exception:
        return False


def formatar_duracao(segundos):
    segundos = int(segundos)
    h, resto = divmod(segundos, 3600)
    m, s = divmod(resto, 60)
    if h:
        return f"{h}h{m:02d}m"
    if m:
        return f"{m}m{s:02d}s"
    return f"{s}s"


def cabecalho_planilha(ws, titulo, cabecalhos, larguras):
    """Título mesclado na linha 1 + cabeçalhos formatados na linha 2."""
    ultima = get_column_letter(len(cabecalhos))
    ws.merge_cells(f"A1:{ultima}1")
    ws["A1"] = titulo
    ws["A1"].font = FONTE_TITULO
    ws["A1"].fill = FILL_TITULO
    ws["A1"].alignment = CENTRO
    ws.row_dimensions[1].height = 30
    for col, cab in enumerate(cabecalhos, 1):
        cell = ws.cell(row=2, column=col, value=cab)
        cell.font = FONTE_CABEC
        cell.fill = FILL_CABEC
        cell.alignment = CENTRO
        cell.border = BORDA
    ws.row_dimensions[2].height = 22
    for col, larg in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(col)].width = larg


def escrever_linha(ws, row_idx, valores, fill, col_texto_esquerda):
    for col_idx, val in enumerate(valores, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = FONTE_NORMAL
        cell.fill = fill
        cell.alignment = Alignment(
            horizontal="left" if col_idx == col_texto_esquerda else "center",
            vertical="center")
        cell.border = BORDA


# ---------------- detecção das colunas pelo cabeçalho ----------------
def _normalizar_texto(valor):
    """'Nº Ocorrência:' → 'n ocorrencia' (sem acento, minúsculo, sem símbolos)."""
    texto = unicodedata.normalize("NFKD", limpar(valor))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"[º°ª.:#/_\-()]", " ", texto.lower())
    return re.sub(r"\s+", " ", texto).strip()


def _pontuar_cabecalho(texto, apelidos):
    """3 = nome idêntico a um apelido | 2 = contém o apelido | 0 = não serve."""
    if not texto:
        return 0
    palavras = texto.split()
    if any(_normalizar_texto(p) in palavras for p in PALAVRAS_IGNORADAS):
        return 0
    apelidos = [_normalizar_texto(a) for a in apelidos]
    if texto in apelidos:
        return 3
    for a in apelidos:
        if re.search(rf"\b{re.escape(a)}\b", texto):
            return 2
    return 0


def detectar_colunas(ws):
    """
    Procura os cabeçalhos de Nota e Ocorrência nas primeiras linhas.
    Devolve (col_nota, col_ocor, linha_inicio, achou_pelo_nome).
    Se não achar, devolve os valores fixos da configuração.
    """
    melhor = {"nota": (0, None, None), "ocor": (0, None, None)}  # (pontos, coluna, linha)
    ultima = min(ws.max_row, LINHAS_BUSCA_CABECALHO)

    for linha in range(1, ultima + 1):
        for cell in ws[linha]:
            if cell.value is None:
                continue
            texto = _normalizar_texto(cell.value)
            for chave, apelidos in (("nota", NOMES_COL_NOTA), ("ocor", NOMES_COL_OCOR)):
                pontos = _pontuar_cabecalho(texto, apelidos)
                if pontos > melhor[chave][0]:
                    melhor[chave] = (pontos, get_column_letter(cell.column), linha)

    _, col_nota, lin_nota = melhor["nota"]
    _, col_ocor, lin_ocor = melhor["ocor"]

    if col_nota and col_ocor and col_nota != col_ocor:
        return col_nota, col_ocor, max(lin_nota, lin_ocor) + 1, True
    return COL_NOTA, COL_OCOR, LINHA_INICIO, False


def _exemplo_coluna(ws, coluna, linha_inicio):
    """Primeiro valor com dígitos da coluna — só para você conferir no terminal."""
    for i in range(linha_inicio, min(ws.max_row, linha_inicio + 50) + 1):
        valor = limpar(ws[f"{coluna}{i}"].value)
        if tem_digitos(valor):
            return valor
    return "(nenhum valor numérico encontrado)"


_COLUNAS_CONFIRMADAS = {}   # (caminho, data de modificação) → (col_nota, col_ocor, inicio)


def obter_colunas(ws, caminho):
    """
    Detecta as colunas, mostra o resultado e pede confirmação UMA vez por
    planilha na sessão. Se o arquivo for salvo de novo (colunas mexidas),
    pergunta de novo.
    """
    try:
        chave = (os.path.abspath(caminho), os.path.getmtime(caminho))
    except OSError:
        chave = (os.path.abspath(caminho), 0)
    if chave in _COLUNAS_CONFIRMADAS:
        return _COLUNAS_CONFIRMADAS[chave]

    col_nota, col_ocor, inicio, achou = detectar_colunas(ws)
    manual = False

    while True:
        cab = inicio - 1
        titulo_nota = limpar(ws[f"{col_nota}{cab}"].value) if cab >= 1 else ""
        titulo_ocor = limpar(ws[f"{col_ocor}{cab}"].value) if cab >= 1 else ""

        print("\n" + "-" * 64)
        if manual:
            print("  Colunas definidas por você:")
        elif achou:
            print("  Colunas encontradas pelo cabeçalho:")
        else:
            print("  ⚠ Não achei os cabeçalhos pelo nome — usando as colunas da configuração:")
        print(f"    Nota       → coluna {col_nota}  ('{titulo_nota}')  "
              f"ex.: {_exemplo_coluna(ws, col_nota, inicio)}")
        print(f"    Ocorrência → coluna {col_ocor}  ('{titulo_ocor}')  "
              f"ex.: {_exemplo_coluna(ws, col_ocor, inicio)}")
        print(f"    Dados a partir da linha {inicio}")
        print("-" * 64)

        resp = input("  ENTER confirma | ou digite NOTA OCORRÊNCIA [linha inicial] "
                     "(ex.: D H  ou  C G 3): ").strip().upper()
        if not resp:
            break

        partes = resp.replace(",", " ").split()
        try:
            nova_nota, nova_ocor = partes[0], partes[1]
            column_index_from_string(nova_nota)
            column_index_from_string(nova_ocor)
            novo_inicio = int(partes[2]) if len(partes) > 2 else inicio
            if nova_nota == nova_ocor or novo_inicio < 1:
                raise ValueError
        except (IndexError, ValueError):
            print("  Entrada inválida. Exemplo: D H   (nota na D, ocorrência na H)")
            continue
        col_nota, col_ocor, inicio, manual = nova_nota, nova_ocor, novo_inicio, True

    _COLUNAS_CONFIRMADAS[chave] = (col_nota, col_ocor, inicio)
    return col_nota, col_ocor, inicio


def ler_planilha(caminho):
    """
    Lê a planilha de entrada e devolve uma lista de registros, cada um com
    um NOME DE ARQUIVO ÚNICO ('12345', '12345 (2)', ...).

    É a ÚNICA leitura da planilha no script inteiro: SAP, verificação e
    reprocessamento usam exatamente os mesmos nomes.
    """
    if not os.path.exists(caminho):
        print(f"ERRO: planilha de entrada não encontrada em:\n  {caminho}")
        return []

    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb[ABA] if ABA else wb.active
    col_nota, col_ocor, linha_inicio = obter_colunas(ws, caminho)

    registros, ignoradas = [], []
    contagem_nota = defaultdict(int)

    for i in range(linha_inicio, ws.max_row + 1):
        ocorrencia_bruta = limpar(ws[f"{col_ocor}{i}"].value)
        num_nota         = limpar(ws[f"{col_nota}{i}"].value)

        if not ocorrencia_bruta:
            if num_nota:
                ignoradas.append((i, "", num_nota))
            continue

        if not tem_digitos(ocorrencia_bruta) or not tem_digitos(num_nota):
            ignoradas.append((i, ocorrencia_bruta, num_nota))
            continue

        contagem_nota[num_nota] += 1
        vez = contagem_nota[num_nota]

        registros.append({
            "linha": i,
            "ocorrencia": extrair_ocorrencia(ocorrencia_bruta),
            "ocorrencia_bruta": ocorrencia_bruta,
            "nota": num_nota,
            "nome_unico": num_nota if vez == 1 else f"{num_nota} ({vez})",
            "duplicada": vez > 1,
        })

    duplicadas = sum(1 for r in registros if r["duplicada"])
    print(f"{len(registros)} linha(s) lidas de {os.path.basename(caminho)} "
          f"({duplicadas} com número de nota repetido).")

    if ignoradas:
        print(f"{len(ignoradas)} linha(s) ignorada(s) (sem ocorrência/nota numérica):")
        for i, ocor, nota in ignoradas:
            print(f"    Linha {i}: ocorrência='{ocor or '(vazio)'}' | nota='{nota or '(vazio)'}'")

    return registros


# =====================================================================
#                ETAPA 1 — SAP (NIXNFSE) — antigo main.py
# =====================================================================
def conectar_sap():
    sap_gui = win32com.client.GetObject("SAPGUI")
    app = sap_gui.GetScriptingEngine
    connection = app.Children(0)
    return connection.Children(0)


def esperar_pdf_pronto(caminho, timeout=TIMEOUT_PDF, intervalo=INTERVALO_POLL):
    """Espera ativa: sai assim que o PDF existir, for válido e com tamanho estável."""
    fim = time.time() + timeout
    tam_anterior, estavel = -1, 0
    while time.time() < fim:
        if pdf_valido(caminho):
            try:
                tam = os.path.getsize(caminho)
            except Exception:
                tam = -1
            if tam == tam_anterior and tam > 0:
                estavel += 1
                if estavel >= 2:
                    return True
            else:
                estavel = 0
                tam_anterior = tam
        time.sleep(intervalo)
    return False


# ---------------- captura do PDF temporário ----------------
def _candidatas_padrao():
    local  = os.environ.get("LOCALAPPDATA", "")
    roam   = os.environ.get("APPDATA", "")
    perfil = os.environ.get("USERPROFILE", "")
    return [
        tempfile.gettempdir(),
        os.path.join(local, "Temp"),
        os.path.join(local, "SAP"),
        os.path.join(roam,  "SAP"),
        os.path.join(perfil, "Documents", "SAP"),
        os.path.join(perfil, "SapWorkDir"),
        os.path.join(local, "Microsoft", "Windows", "INetCache"),
    ]


def _varrer_pasta(raiz):
    """{caminho: (mtime, tamanho)} com profundidade limitada (os.scandir)."""
    encontrados = {}
    pilha = [(raiz, 0)]
    while pilha:
        atual, nivel = pilha.pop()
        try:
            with os.scandir(atual) as it:
                for entrada in it:
                    try:
                        if entrada.is_dir(follow_symlinks=False):
                            if nivel < PROFUNDIDADE_VARREDURA:
                                pilha.append((entrada.path, nivel + 1))
                        elif entrada.is_file(follow_symlinks=False):
                            st = entrada.stat()
                            if st.st_size <= TAMANHO_MAXIMO_PDF:
                                encontrados[entrada.path] = (st.st_mtime, st.st_size)
                    except OSError:
                        continue
        except (PermissionError, OSError):
            continue
    return encontrados


_PASTAS = None
_PASTA_TRAVADA = False
_ORIGENS_USADAS = {}   # impede reaproveitar o mesmo temporário em outra nota


def calibrar_pastas():
    global _PASTAS
    if _PASTA_TRAVADA and _PASTAS:
        return _PASTAS   # já descobriu a pasta certa numa rodada anterior da sessão

    brutas = list(PASTAS_MONITORADAS) if PASTAS_MONITORADAS else _candidatas_padrao()
    vistas = []
    for c in brutas:
        if not c:
            continue
        c = os.path.normpath(c)
        if os.path.isdir(c) and c not in vistas:
            vistas.append(c)

    print("\n  Calibrando pastas de captura...")
    aprovadas = []
    for pasta in vistas:
        t0 = time.time()
        arquivos = _varrer_pasta(pasta)
        gasto = time.time() - t0
        if gasto > LIMITE_TEMPO_VARREDURA:
            print(f"    [ignorada] {gasto:5.2f}s  {len(arquivos):>6} arq  {pasta}")
            print(f"               (lenta demais — passaria {gasto*2:.1f}s por nota)")
        else:
            print(f"    [ok]       {gasto:5.2f}s  {len(arquivos):>6} arq  {pasta}")
            aprovadas.append(pasta)

    _PASTAS = aprovadas
    if not _PASTAS:
        print("\n  AVISO: nenhuma pasta aprovada. A captura vai falhar e o script")
        print("  vai depender do método antigo (tela ocupada).")
    return _PASTAS


def tirar_foto():
    if _PASTAS is None:
        calibrar_pastas()
    encontrados = {}
    for raiz in _PASTAS:
        encontrados.update(_varrer_pasta(raiz))
    return encontrados


def _copiar_com_retry(origem, destino, tentativas=5, espera=0.3):
    for _ in range(tentativas):
        try:
            shutil.copy2(origem, destino)
            return True
        except (PermissionError, OSError):
            time.sleep(espera)
    return False


def capturar_pdf_novo(foto_antes, caminho_destino,
                      timeout=TIMEOUT_CAPTURA, intervalo=INTERVALO_CAPTURA):
    """Espera surgir um PDF novo no temp (estável, válido, não reutilizado) e copia."""
    inicio = time.time()
    fim = inicio + timeout
    corte_mtime = inicio - 2.0
    leituras = {}
    novos_ignorados = {}

    while time.time() < fim:
        candidatos = []
        for caminho, (mtime, tamanho) in tirar_foto().items():
            anterior = foto_antes.get(caminho)
            if anterior is not None and anterior == (mtime, tamanho):
                continue
            if mtime < corte_mtime:
                continue
            if _ORIGENS_USADAS.get(caminho) == (mtime, tamanho):
                continue
            if tamanho < TAMANHO_MINIMO_PDF or not pdf_valido(caminho):
                if tamanho > 0:
                    novos_ignorados[caminho] = tamanho
                leituras.pop(caminho, None)
                continue

            tam_anterior, repeticoes = leituras.get(caminho, (None, 0))
            repeticoes = repeticoes + 1 if tam_anterior == tamanho else 0
            leituras[caminho] = (tamanho, repeticoes)
            if repeticoes >= LEITURAS_ESTAVEIS:
                candidatos.append((mtime, caminho, tamanho))

        if candidatos:
            candidatos.sort(reverse=True)
            _, origem, tamanho = candidatos[0]
            if _copiar_com_retry(origem, caminho_destino):
                try:
                    _ORIGENS_USADAS[origem] = (os.path.getmtime(origem), tamanho)
                except OSError:
                    _ORIGENS_USADAS[origem] = (0, tamanho)
                if not _PASTA_TRAVADA:
                    _travar_na_pasta(origem)
                return True

        time.sleep(intervalo)

    _relatar_falha_captura(novos_ignorados)
    return False


def _travar_na_pasta(origem):
    global _PASTAS, _PASTA_TRAVADA
    vencedora = next((raiz for raiz in _PASTAS if origem.startswith(raiz)), None)
    if vencedora:
        _PASTAS = [vencedora]
        _PASTA_TRAVADA = True
        print(f"\n  >> Captura confirmada. Pasta usada pelo SAP:")
        print(f"     {os.path.dirname(origem)}")
        print(f"     A varredura agora fica restrita a: {vencedora}")
        print(f"     (para acelerar de vez, fixe isso em PASTAS_MONITORADAS)\n")


def _relatar_falha_captura(novos_ignorados):
    print(f"    (debug) nenhum PDF novo em {TIMEOUT_CAPTURA:.0f}s. Pastas vigiadas: {_PASTAS}")
    if novos_ignorados:
        print(f"    (debug) {len(novos_ignorados)} arquivo(s) novo(s) apareceram, "
              f"mas nenhum tem assinatura de PDF. Os maiores:")
        for caminho, tamanho in sorted(novos_ignorados.items(), key=lambda x: -x[1])[:5]:
            print(f"              {tamanho:>10} bytes  {caminho}")
    else:
        print("    (debug) nenhum arquivo novo apareceu — provável sinal de que o "
              "SAP renderiza em memória, ou de que a pasta certa ficou de fora.")


# ---------------- método antigo (fallback) ----------------
def selecionar_tipo_pdf(dlg):
    try:
        combos = dlg.children(class_name="ComboBox") or dlg.children(class_name="ComboBoxEx32")
        for combo in combos:
            try:
                itens = combo.item_texts()
            except Exception:
                continue
            for alvo in TIPOS_PDF_ACEITOS:
                for item in itens:
                    if alvo.lower() in item.lower():
                        combo.select(item)
                        time.sleep(0.2)
                        return True
    except Exception as e:
        print(f"  aviso: não foi possível checar/selecionar tipo de arquivo: {e}")
    return False


def salvar_via_pywinauto(caminho_completo):
    try:
        desktop = Desktop(backend="win32")
        dlg = desktop.window(title_re="Salvar como|Save As|Salvar Saída de Impressão")
        dlg.wait("visible", timeout=10)
        time.sleep(0.15)
        selecionar_tipo_pdf(dlg)
        try:
            dlg.child_window(class_name="Edit").set_edit_text(caminho_completo)
        except Exception:
            edits = dlg.children(class_name="Edit")
            if edits:
                edits[-1].set_edit_text(caminho_completo)
        time.sleep(0.2)
        try:
            dlg.child_window(title_re="&Salvar|Salvar|Save", class_name="Button").click()
        except Exception:
            dlg.type_keys("{ENTER}")
        return True
    except Exception as e:
        print(f"  pywinauto erro: {e} — tentando fallback pyautogui...")
        pyperclip.copy(caminho_completo)
        pyautogui.hotkey("ctrl", "a")
        time.sleep(0.2)
        pyautogui.hotkey("ctrl", "v")
        time.sleep(0.3)
        pyautogui.press("enter")
        return False


def salvar_nfse_fallback(session, caminho_completo):
    """MÉTODO ANTIGO — exige a tela livre e o SAP em primeiro plano."""
    pyautogui.click(*TRES_PONTOS)
    time.sleep(ESPERA_MENU)
    for _ in range(SETAS_ATE_SALVAR):
        pyautogui.press("down")
        time.sleep(0.2)
    pyautogui.press("enter")
    time.sleep(0.4)
    salvar_via_pywinauto(caminho_completo)


# ---------------- navegação no NIXNFSE ----------------
def abrir_selecao(session):
    session.findById("wnd[0]/tbar[0]/okcd").Text = "/nNIXNFSE"
    session.findById("wnd[0]").sendVKey(0)
    session.findById(
        "wnd[0]/usr/cntlIMAGE_CONTAINER/shellcont/shell/shellcont[0]/shell"
    ).doubleClickNode("F00003")


def voltar_para_selecao(session):
    for _ in range(VOLTAS_ATE_SELECAO):
        try:
            session.findById("wnd[0]").sendVKey(3)
            time.sleep(0.3)
        except Exception:
            break
    try:
        session.findById("wnd[0]/usr/ssubSUB_SEL:/MIDAS/NIXNFSE_COCKPIT:0001/txtS_OCOR-LOW")
        return True
    except Exception:
        return False


def limpar_extras(pasta, nome_base):
    pasta_files = os.path.join(pasta, f"{nome_base}_files")
    if os.path.isdir(pasta_files):
        shutil.rmtree(pasta_files, ignore_errors=True)
    try:
        for nome in os.listdir(pasta):
            base, ext = os.path.splitext(nome)
            if base == nome_base and ext.lower() != ".pdf":
                try:
                    os.remove(os.path.join(pasta, nome))
                except Exception:
                    pass
    except Exception:
        pass


def imprimir_parcial(processadas, total, ok, erros, materiais, inicio):
    decorrido = time.time() - inicio
    media = decorrido / processadas if processadas else 0
    eta_seg = media * (total - processadas)
    hora_fim = (datetime.now() + timedelta(seconds=eta_seg)).strftime("%H:%M")
    pct = (processadas / total * 100) if total else 0
    print(f"  ── PARCIAL {processadas}/{total} ({pct:.0f}%) | "
          f"✔ {ok} OK | ✘ {erros} erro(s) | {materiais} material | "
          f"média {media:.1f}s/nota | restam ~{formatar_duracao(eta_seg)} "
          f"(término ~{hora_fim}) ──")


def baixar_notas_sap(registros):
    """Baixa as notas de serviço do SAP para cada registro. Devolve os dados do relatório."""
    _carregar_libs_sap()
    os.makedirs(PASTA_DESTINO, exist_ok=True)

    total = len(registros)
    if LIMITE_NOTAS and total > LIMITE_NOTAS:
        registros = registros[:LIMITE_NOTAS]
        total = len(registros)
        print(f"\n  *** MODO TESTE: processando só as {LIMITE_NOTAS} primeiras "
              f"notas (LIMITE_NOTAS). Coloque 0 para rodar tudo. ***")

    print("=" * 64)
    print(f"  SAP — INÍCIO {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    print(f"  Destino  : {PASTA_DESTINO}")
    print(f"  Notas a processar: {total}")
    print("=" * 64)

    calibrar_pastas()
    session = conectar_sap()
    if USAR_FALLBACK_PYAUTOGUI:
        session.findById("wnd[0]").maximize()

    dados_relatorio = []
    selecao_pronta = False
    inicio_geral = time.time()
    processadas = cont_ok = cont_erro = cont_material = 0

    for r in registros:
        i                = r["linha"]
        ocorrencia       = r["ocorrencia"]
        ocorrencia_bruta = r["ocorrencia_bruta"]
        num_nota         = r["nota"]
        nome_base        = r["nome_unico"]
        nome_arquivo     = f"{nome_base}.pdf"
        caminho          = os.path.join(PASTA_DESTINO, nome_arquivo)

        processadas += 1
        inicio_nota = time.time()
        prefixo = f"[{processadas}/{total}] Linha {i}"
        print(f"{prefixo} | Ocorrência {ocorrencia} | Arquivo {nome_arquivo} → processando...")

        base_rel = {"linha": i, "ocorrencia": ocorrencia, "nota": num_nota,
                    "arquivo": nome_base}

        try:
            if REUSAR_TELA_SELECAO and selecao_pronta:
                if not voltar_para_selecao(session):
                    abrir_selecao(session)
            else:
                abrir_selecao(session)
            selecao_pronta = True

            session.findById(
                "wnd[0]/usr/ssubSUB_SEL:/MIDAS/NIXNFSE_COCKPIT:0001/ctxtS_DTINTE-LOW").Text = ""
            session.findById(
                "wnd[0]/usr/ssubSUB_SEL:/MIDAS/NIXNFSE_COCKPIT:0001/txtS_OCOR-LOW").Text = ocorrencia
            session.findById("wnd[0]/tbar[1]/btn[8]").press()

            try:
                grid = session.findById("wnd[0]/usr/cntlCONT_COCKPIT/shellcont/shell")
                row_count = grid.RowCount
            except Exception:
                row_count = 0

            if row_count == 0:
                dados_relatorio.append({**base_rel, "tipo": "Material",
                                        "status": "Não encontrada no SAP",
                                        "obs": "NFS-e não localizada para esta ocorrência"})
                cont_material += 1
                print(f"{prefixo} → MATERIAL (não encontrada no SAP) "
                      f"[{time.time() - inicio_nota:.1f}s]")
                if REUSAR_TELA_SELECAO:
                    selecao_pronta = False
                if processadas % PARCIAL_A_CADA == 0:
                    imprimir_parcial(processadas, total, cont_ok, cont_erro,
                                     cont_material, inicio_geral)
                continue

            print(f"{prefixo} → encontrada no SAP, abrindo PDF...")
            grid.currentCellColumn = "OCORRENCIA"
            grid.selectedRows = "0"
            grid.clickCurrentCell()
            try:
                session.findById("wnd[1]").sendVKey(0)
            except Exception:
                pass

            # a foto TEM que ser tirada antes de mandar abrir a nota
            foto = tirar_foto()
            session.findById("wnd[0]/tbar[1]/btn[7]").press()

            print(f"{prefixo} → capturando {nome_arquivo}...")
            capturado = capturar_pdf_novo(foto, caminho)

            if not capturado and USAR_FALLBACK_PYAUTOGUI:
                print(f"{prefixo} → captura falhou, usando o método antigo "
                      f"(NÃO MEXA NO COMPUTADOR AGORA)...")
                time.sleep(ESPERA_RENDER)
                salvar_nfse_fallback(session, caminho)
                esperar_pdf_pronto(caminho)

            limpar_extras(PASTA_DESTINO, nome_base)

            if pdf_valido(caminho):
                status, obs = "OK", caminho
                if ocorrencia != ocorrencia_bruta:
                    obs += f"  (ocorrência original: {ocorrencia_bruta} → usada: {ocorrencia})"
                cont_ok += 1
                print(f"{prefixo} → ✔ {nome_arquivo} salvo [{time.time() - inicio_nota:.1f}s]")
            else:
                if os.path.exists(caminho):
                    try:
                        os.remove(caminho)
                    except Exception:
                        pass
                status = "ERRO: PDF não capturado"
                obs = ("A nota abriu no SAP mas nenhum PDF novo apareceu nas pastas "
                       "monitoradas dentro do TIMEOUT_CAPTURA.")
                cont_erro += 1
                print(f"{prefixo} → ✘ FALHOU: PDF inválido [{time.time() - inicio_nota:.1f}s]")

        except Exception as e:
            status, obs = f"ERRO: {e}", str(e)
            cont_erro += 1
            print(f"{prefixo} → ✘ ERRO: {e} [{time.time() - inicio_nota:.1f}s]")
            selecao_pronta = False

        dados_relatorio.append({**base_rel, "tipo": "Serviço", "status": status, "obs": obs})

        if processadas % PARCIAL_A_CADA == 0:
            imprimir_parcial(processadas, total, cont_ok, cont_erro,
                             cont_material, inicio_geral)

    duracao = time.time() - inicio_geral
    if total:
        print(f"\nSAP concluído em {formatar_duracao(duracao)} "
              f"({duracao / total:.1f}s por nota em média).")
    return dados_relatorio


# Relatório do SAP: a coluna "Arquivo" leva o nome único, e é por ela que o
# NIX sabe com que nome salvar o PDF (ex.: '12345 (2).pdf').
CABECALHOS_RELATORIO = ["Linha", "Ocorrência", "Nota Fiscal", "Arquivo",
                        "Tipo", "Status PDF", "Observação"]
COL_REL_STATUS = 6
COL_REL_OBS = 7


def gerar_relatorio_excel(dados):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relatório NFS-e"
    cabecalho_planilha(ws,
                       f"Relatório NFS-e  —  Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                       CABECALHOS_RELATORIO, [8, 18, 18, 20, 12, 22, 35])

    for row_idx, d in enumerate(dados, 3):
        if d["tipo"] == "Serviço" and d["status"] == "OK":
            fill = FILL_VERDE
        elif d["tipo"] == "Serviço":
            fill = FILL_AMARELO
        else:
            fill = FILL_VERMELHO
        valores = [d["linha"], d["ocorrencia"], d["nota"], d["arquivo"],
                   d["tipo"], d["status"], d.get("obs", "")]
        escrever_linha(ws, row_idx, valores, fill, col_texto_esquerda=COL_REL_OBS)

    try:
        wb.save(RELATORIO_EXCEL)
    except PermissionError:
        print("\n⚠️  FECHE o _relatorio_nfse.xlsx no Excel — não foi possível salvá-lo.")
        return None
    return RELATORIO_EXCEL


def imprimir_relatorio_sap(dados):
    servicos  = [d for d in dados if d["tipo"] == "Serviço"]
    materiais = [d for d in dados if d["tipo"] == "Material"]
    ok        = [d for d in servicos if d["status"] == "OK"]
    erros     = [d for d in servicos if d["status"] != "OK"]

    print("\n" + "=" * 60)
    print("           RELATÓRIO FINAL — SAP (NFS-e)")
    print("=" * 60)
    print(f"  Total de ocorrências           : {len(dados)}")
    print(f"  Serviço  (NFS-e encontrada)    : {len(servicos)}")
    print(f"    ✔  PDFs salvos com sucesso   : {len(ok)}")
    print(f"    ✘  Erros ao salvar PDF       : {len(erros)}")
    print(f"  Material (NFS-e não encontrada): {len(materiais)}")
    print("=" * 60)

    if materiais:
        print("\n  Ocorrências → MATERIAL (vão para o NIX web):")
        for d in materiais:
            print(f"    Linha {d['linha']:>4} | Ocorrência {d['ocorrencia']} | {d['arquivo']}.pdf")
    if erros:
        print("\n  Serviços com ERRO ao salvar PDF:")
        for d in erros:
            print(f"    Linha {d['linha']:>4} | Ocorrência {d['ocorrencia']} | {d['status']}")
    print("=" * 60 + "\n")


# =====================================================================
#            ETAPA 2 — NIX WEB (Material) — antigo nix_material.py
# =====================================================================
_INDICE_COLUNA_CACHE = None


def normalizar_ocorrencia(valor):
    digitos = re.sub(r"\D", "", str(valor))
    return digitos if digitos else None


def ler_materiais():
    """Lê as linhas 'Material' do relatório gerado pela etapa do SAP."""
    if not os.path.exists(RELATORIO_EXCEL):
        print(f"ERRO: relatório não encontrado em:\n  {RELATORIO_EXCEL}")
        print("Rode a etapa do SAP (opção 1) primeiro.")
        return []

    wb = openpyxl.load_workbook(RELATORIO_EXCEL, data_only=True)
    ws = wb.active
    formato_novo = ws.cell(row=2, column=4).value == "Arquivo"

    materiais, ignoradas = [], 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or not row[0]:
            continue
        if formato_novo:
            linha, ocorrencia, nota, arquivo, tipo = row[:5]
        else:  # relatório gerado pelo main.py antigo (sem coluna Arquivo)
            linha, ocorrencia, nota, tipo = row[:4]
            arquivo = nota
        if str(tipo).strip().lower() != "material":
            continue
        ocorrencia_num = normalizar_ocorrencia(ocorrencia)
        if not ocorrencia_num:
            ignoradas += 1
            print(f"  ⚠ Linha {linha}: ocorrência '{ocorrencia}' não é numérica — pulando.")
            continue
        materiais.append({"linha": linha, "ocorrencia": ocorrencia_num,
                          "nota": limpar(nota), "arquivo": limpar(arquivo)})

    print(f"{len(materiais)} nota(s) de Material encontradas no relatório.")
    if ignoradas:
        print(f"{ignoradas} linha(s) ignorada(s) por não terem ocorrência numérica.")
    return materiais


def atualizar_relatorio(resultados):
    try:
        wb = openpyxl.load_workbook(RELATORIO_EXCEL)
    except Exception as e:
        print(f"\n⚠️  Não consegui abrir o _relatorio_nfse.xlsx.\n{e}")
        return

    ws = wb.active
    if ws.cell(row=2, column=4).value == "Arquivo":
        col_status, col_obs, ultima_col = COL_REL_STATUS, COL_REL_OBS, 7
    else:
        col_status, col_obs, ultima_col = 5, 6, 6

    for row_idx in range(3, ws.max_row + 1):
        linha_planilha = ws.cell(row=row_idx, column=1).value
        if linha_planilha in resultados:
            r = resultados[linha_planilha]
            fill = FILL_VERDE if r["status"] == "OK" else FILL_AMARELO
            for col in range(1, ultima_col + 1):
                cell = ws.cell(row=row_idx, column=col)
                cell.fill = fill
                cell.font = FONTE_NORMAL
                cell.border = BORDA
            ws.cell(row=row_idx, column=col_status).value = r["status"]
            ws.cell(row=row_idx, column=col_obs).value = r["obs"]

    try:
        wb.save(RELATORIO_EXCEL)
        print("Relatório atualizado com sucesso.")
    except PermissionError:
        print("\n⚠️  Feche o _relatorio_nfse.xlsx e tente novamente.")


def encontrar_botao_download(pagina, timeout_ms):
    """Procura o #download do PDF.js na página e, se não achar, dentro dos iframes."""
    try:
        btn = pagina.locator("#download")
        btn.wait_for(state="visible", timeout=2000)
        return btn
    except PlaywrightTimeout:
        pass

    deadline = time.time() + (timeout_ms / 1000)
    while time.time() < deadline:
        for frame in pagina.frames:
            try:
                btn = frame.locator("#download")
                if btn.count() > 0:
                    btn.wait_for(state="visible", timeout=1000)
                    return btn
            except Exception:
                continue
        pagina.wait_for_timeout(300)

    try:
        print(f"    (debug) frames encontrados na página: {[f.url for f in pagina.frames]}")
    except Exception:
        pass
    return None


def salvar_diagnostico_visual(pagina, nome_base):
    os.makedirs(PASTA_DIAGNOSTICO, exist_ok=True)
    caminho_print = os.path.join(PASTA_DIAGNOSTICO, f"{nome_base}_falha.png")
    try:
        pagina.screenshot(path=caminho_print, timeout=5000)
    except Exception:
        caminho_print = None
    return caminho_print


def _indice_coluna(linha, nome_coluna):
    """Posição da coluna no grid DevExpress (cabeçalhos são td.dxgvHeader, não <th>)."""
    containers = [
        "xpath=ancestor::table[1]",
        "xpath=ancestor::table[last()]",
        "xpath=ancestor::div[contains(@class,'dxgv')][last()]",
        "xpath=ancestor::div[last()]",
    ]
    seletores_header = [
        ".//td[contains(@class,'dxgvHeader')]",
        ".//th[contains(@class,'dxgvHeader')]",
        ".//*[self::td or self::th][contains(@class,'Header')]",
        ".//thead//th",
        ".//tr[th][1]/th",
    ]
    alvo = re.sub(r"\s+", " ", nome_coluna).strip().lower()

    try:
        total_td = linha.locator("xpath=./td").count()
    except Exception:
        return None

    textos_vistos = []
    for container_xpath in containers:
        try:
            container = linha.locator(container_xpath)
            if container.count() == 0:
                continue
        except Exception:
            continue

        for seletor in seletores_header:
            try:
                cabecalhos = container.locator(f"xpath={seletor}")
                total_th = cabecalhos.count()
            except Exception:
                continue
            if total_th == 0:
                continue

            for i in range(total_th):
                try:
                    texto = cabecalhos.nth(i).inner_text(timeout=2000)
                except Exception:
                    continue
                texto = re.sub(r"\s+", " ", texto).strip().lower()
                if texto:
                    textos_vistos.append(texto)
                if alvo in texto:
                    return i + max(total_td - total_th, 0)

    if textos_vistos:
        print(f"    (debug) cabecalhos lidos no grid: {sorted(set(textos_vistos))}")
    return None


def _digitos_celula(celulas, indice, timeout=3000):
    try:
        return re.sub(r"\D", "", celulas.nth(indice).inner_text(timeout=timeout))
    except Exception:
        return None


def verificar_resultado_correto(botao_visualizar, ocorrencia_esperada, silencioso=False):
    """Confere, por igualdade exata dos dígitos, se a linha é a ocorrência buscada."""
    global _INDICE_COLUNA_CACHE

    try:
        linha = botao_visualizar.locator("xpath=ancestor::tr[1]")
        celulas = linha.locator("xpath=./td")
        total_td = celulas.count()
    except Exception:
        return False
    if total_td == 0:
        return False

    idx = _INDICE_COLUNA_CACHE
    if idx is None:
        idx = _indice_coluna(linha, COLUNA_OCORRENCIA)
        if idx is not None:
            _INDICE_COLUNA_CACHE = idx

    if idx is not None and idx < total_td:
        valor = _digitos_celula(celulas, idx)
        if valor == ocorrencia_esperada:
            return True
        if not silencioso:
            print(f"    (debug) coluna '{COLUNA_OCORRENCIA}' (posicao {idx}) da linha "
                  f"exibida = '{valor or '(vazio)'}' | esperado = '{ocorrencia_esperada}'")
        return False

    if not silencioso:
        print(f"    (aviso) coluna '{COLUNA_OCORRENCIA}' nao localizada no cabecalho; "
              f"conferindo celula a celula.")
    for i in range(total_td):
        valor = _digitos_celula(celulas, i, timeout=2000)
        if valor and valor == ocorrencia_esperada:
            return True
    return False


def aguardar_linha_da_ocorrencia(page, ocorrencia_esperada, timeout_ms):
    """Espera o grid mostrar a linha da ocorrência buscada (a busca é AJAX)."""
    deadline = time.time() + (timeout_ms / 1000)
    while time.time() < deadline:
        botoes = page.locator("a.MvcGridSearchButton")
        try:
            total = botoes.count()
        except Exception:
            total = 0
        for i in range(total):
            btn = botoes.nth(i)
            try:
                if not btn.is_visible():
                    continue
            except Exception:
                continue
            if verificar_resultado_correto(btn, ocorrencia_esperada, silencioso=True):
                return btn
        page.wait_for_timeout(300)
    return None


def descrever_linha_exibida(page, ocorrencia_esperada):
    try:
        botoes = page.locator("a.MvcGridSearchButton")
        if botoes.count() == 0:
            print("    (debug) nenhum resultado no grid.")
            return
        verificar_resultado_correto(botoes.first, ocorrencia_esperada, silencioso=False)
    except Exception:
        pass


def baixar_pdf_via_botao_viewer(page, context, botao_visualizar, caminho_destino,
                                timeout_ms=TIMEOUT_RESPOSTA_PDF):
    pagina_viewer = page
    try:
        with context.expect_page(timeout=5000) as new_page_info:
            botao_visualizar.click()
        pagina_viewer = new_page_info.value
        pagina_viewer.wait_for_load_state("domcontentloaded", timeout=timeout_ms)
    except PlaywrightTimeout:
        try:
            page.locator("#divModalTitle").wait_for(state="visible", timeout=5000)
        except Exception:
            pass
        pagina_viewer = page

    try:
        btn_download = encontrar_botao_download(pagina_viewer, timeout_ms)
        if btn_download is None:
            caminho_print = salvar_diagnostico_visual(
                pagina_viewer, os.path.splitext(os.path.basename(caminho_destino))[0])
            raise Exception(
                "Botão de download do visualizador não encontrado (nem na página "
                "nem em iframes)." + (f" Screenshot salvo em: {caminho_print}" if caminho_print else ""))

        with pagina_viewer.expect_download(timeout=timeout_ms) as download_info:
            btn_download.click()
        download_info.value.save_as(caminho_destino)
        return True
    finally:
        if pagina_viewer is not page:
            try:
                pagina_viewer.close()
            except Exception:
                pass


MAX_QUEDAS_SEGUIDAS = 2   # quantas vezes o navegador pode cair na MESMA nota antes de parar


def _navegador_fechado(erro):
    """True se o erro indica que o navegador/aba morreu (e não um erro da nota)."""
    texto = str(erro).lower()
    return ("has been closed" in texto or "target closed" in texto
            or "browser has disconnected" in texto or "target crashed" in texto)


def abrir_navegador_nix(p):
    """Abre o Edge com o perfil da macro, entra no NIX e trata o login."""
    context = p.chromium.launch_persistent_context(
        PERFIL_NAVEGADOR, headless=False, accept_downloads=True,
        channel="msedge", ignore_https_errors=True,
        args=["--start-maximized"], no_viewport=True)

    page = context.pages[0] if context.pages else context.new_page()
    # se a aba travar, aparece no terminal (diferencia "travou" de "fechou")
    page.on("crash", lambda *_: print("    (debug) a aba do navegador TRAVOU (crash)."))

    page.goto(URL_NIX, wait_until="domcontentloaded")
    time.sleep(2)

    if "Login" in page.title() or "Account" in page.url:
        print("Tela de login detectada — clicando em Entrar...")
        try:
            btn_entrar = page.locator("#btn-modal")
            btn_entrar.wait_for(state="visible", timeout=5000)
            btn_entrar.click()
            time.sleep(2)
            page.wait_for_url(lambda url: "Login" not in url and "Account" not in url,
                              timeout=15000)
            print("Login realizado com sucesso!")
        except Exception as e:
            print(f"Não consegui clicar em Entrar automaticamente: {e}")
            input("Faça o login manualmente e pressione ENTER...")
        page.goto(URL_NIX, wait_until="domcontentloaded")
        time.sleep(2)

    return context, page


def _baixar_um_material(page, context, item, n, total):
    """
    Processa UMA nota de material.
    Devolve (resultado, navegador_caiu). Se navegador_caiu for True, quem
    chamou reabre o navegador e tenta esta mesma nota de novo.
    """
    ocorrencia = item["ocorrencia"]
    linha      = item["linha"]
    nome_pdf   = f"{item['arquivo']}.pdf"
    caminho    = os.path.join(PASTA_DESTINO, nome_pdf)

    print(f"\n[{n}/{total}] Linha {linha}: ocorrência {ocorrencia} "
          f"→ buscando no NIX ({nome_pdf})...")

    try:
        if URL_NIX not in page.url:
            page.goto(URL_NIX, wait_until="domcontentloaded")
            time.sleep(1.5)

        campo_ocor = page.locator("#GaiaOccurenceId")
        campo_ocor.wait_for(state="visible", timeout=10000)
        campo_ocor.fill(ocorrencia)
        for seletor in ("#InitialDateCreate", "#EndDateCreate"):
            campo_data = page.locator(seletor)
            try:
                if campo_data.is_enabled(timeout=2000):
                    campo_data.fill("")
            except Exception:
                pass
        time.sleep(0.3)
        campo_ocor.press("Enter")

        btn_visualizar = aguardar_linha_da_ocorrencia(page, ocorrencia, int(ESPERA_BUSCA * 1000))

        if btn_visualizar is None:
            descrever_linha_exibida(page, ocorrencia)
            caminho_print = salvar_diagnostico_visual(page, f"linha{linha}_ocor{ocorrencia}")
            print(f"  → [FALHOU] Ocorrência {ocorrencia} não apareceu no grid "
                  f"dentro de {ESPERA_BUSCA:.0f}s.")
            return {"status": "ERRO: ocorrência não encontrada no NIX",
                    "obs": "A busca não retornou uma linha com esta ocorrência "
                           "(ou o grid não atualizou a tempo); nada foi baixado para "
                           "evitar salvar a nota errada."
                           + (f" Screenshot: {caminho_print}" if caminho_print else "")}, False

        baixar_pdf_via_botao_viewer(page, context, btn_visualizar, caminho)

        if pdf_valido(caminho):
            print(f"  → {nome_pdf} salvo com sucesso [OK]")
            resultado = {"status": "OK", "obs": caminho}
        else:
            if os.path.exists(caminho):
                os.remove(caminho)
            print("  → [FALHOU] arquivo baixado não é um PDF válido")
            resultado = {"status": "ERRO: arquivo baixado não é PDF válido",
                         "obs": "O download ocorreu mas o conteúdo salvo não tem assinatura de PDF"}

    except Exception as e:
        if _navegador_fechado(e):
            print(f"  → [NAVEGADOR CAIU] {e}")
            return {"status": "ERRO: navegador fechou durante a nota", "obs": str(e)}, True
        print(f"  → [FALHOU] {e}")
        resultado = {"status": f"ERRO: {e}", "obs": str(e)}

    # volta para a tela de busca limpa para a próxima nota
    try:
        page.goto(URL_NIX, wait_until="domcontentloaded")
        time.sleep(1.5)
    except Exception as e:
        if _navegador_fechado(e):
            return resultado, True
    return resultado, False


def processar_materiais_nix(materiais):
    _carregar_libs_nix()
    os.makedirs(PERFIL_NAVEGADOR, exist_ok=True)
    resultados = {}
    total = len(materiais)

    with sync_playwright() as p:
        context, page = abrir_navegador_nix(p)

        idx = 0
        quedas_seguidas = 0
        while idx < total:
            item = materiais[idx]
            resultado, caiu = _baixar_um_material(page, context, item, idx + 1, total)

            if caiu:
                quedas_seguidas += 1
                try:
                    context.close()
                except Exception:
                    pass

                if quedas_seguidas >= MAX_QUEDAS_SEGUIDAS:
                    resultados[item["linha"]] = resultado
                    restantes = total - idx - 1
                    print("\n" + "!" * 64)
                    print(f"  O navegador caiu {quedas_seguidas}x seguidas na MESMA nota "
                          f"(linha {item['linha']}).")
                    print("  Isso indica problema no Edge/Playwright, não na nota.")
                    print(f"  Etapa do NIX interrompida — {restantes} nota(s) ficaram sem processar")
                    print("  e continuam como pendentes (dá para retomar pela opção 5).")
                    print("!" * 64)
                    return resultados

                print("  ↻ Reabrindo o navegador e tentando esta nota de novo...")
                context, page = abrir_navegador_nix(p)
                continue          # mesma nota, sem avançar o idx

            quedas_seguidas = 0
            resultados[item["linha"]] = resultado
            idx += 1

        try:
            context.close()
        except Exception:
            pass

    return resultados


def imprimir_resumo_nix(resultados):
    ok    = [r for r in resultados.values() if r["status"] == "OK"]
    erros = [r for r in resultados.values() if r["status"] != "OK"]
    print("\n" + "=" * 60)
    print("        RELATÓRIO FINAL — NIX Material")
    print("=" * 60)
    print(f"  Total processado      : {len(resultados)}")
    print(f"  ✔  Salvos com sucesso : {len(ok)}")
    print(f"  ✘  Erros              : {len(erros)}")
    print("=" * 60 + "\n")


# =====================================================================
#            ETAPA 3 — VERIFICAÇÃO — antigo verificanota.py
# =====================================================================
def verificar(registros):
    resultados = []
    for r in registros:
        caminho = os.path.join(PASTA_DESTINO, f"{r['nome_unico']}.pdf")
        existe = os.path.exists(caminho)
        valido = pdf_valido(caminho) if existe else False

        if valido:
            situacao, detalhe = "OK", f"{r['nome_unico']}.pdf presente e válido"
        elif existe:
            situacao, detalhe = "PDF CORROMPIDO", f"{r['nome_unico']}.pdf existe mas não é um PDF válido"
        else:
            situacao = "FALTANDO"
            detalhe = f"Esperado '{r['nome_unico']}.pdf' — não está na pasta"
            if r["duplicada"]:
                detalhe += " (nota com número repetido)"

        resultados.append({**r, "situacao": situacao, "detalhe": detalhe})
    return resultados


def listar_pdfs_orfaos(registros):
    esperados = {r["nome_unico"] for r in registros}
    orfaos = []
    try:
        for nome in os.listdir(PASTA_DESTINO):
            base, ext = os.path.splitext(nome)
            if ext.lower() == ".pdf" and base not in esperados:
                orfaos.append(nome)
    except Exception:
        pass
    return sorted(orfaos)


def gerar_planilha_verificacao(resultados, orfaos):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verificação"
    cabecalho_planilha(ws,
                       f"Verificação NFS-e — {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                       ["Linha", "Ocorrência", "Nota Fiscal", "Arquivo esperado",
                        "Duplicada?", "Situação", "Detalhe"],
                       [8, 20, 18, 24, 12, 16, 60])

    for row_idx, r in enumerate(resultados, 3):
        fill = {"OK": FILL_VERDE, "FALTANDO": FILL_AMARELO}.get(r["situacao"], FILL_VERMELHO)
        valores = [r["linha"], r["ocorrencia"], r["nota"], f"{r['nome_unico']}.pdf",
                   "SIM" if r["duplicada"] else "", r["situacao"], r["detalhe"]]
        escrever_linha(ws, row_idx, valores, fill, col_texto_esquerda=7)

    if orfaos:
        ws2 = wb.create_sheet("PDFs fora da lista")
        ws2["A1"] = "Arquivos .pdf presentes na pasta mas sem linha correspondente"
        ws2["A1"].font = Font(bold=True)
        for idx, nome in enumerate(orfaos, 3):
            ws2.cell(row=idx, column=1, value=nome)
        ws2.column_dimensions["A"].width = 40

    try:
        wb.save(ARQUIVO_VERIFICACAO)
    except PermissionError:
        print("\n⚠️  FECHE o _verificacao_nfse.xlsx no Excel — não foi possível salvá-lo.")
        return None
    return ARQUIVO_VERIFICACAO


def imprimir_resumo_verificacao(resultados, orfaos):
    ok          = [r for r in resultados if r["situacao"] == "OK"]
    faltando    = [r for r in resultados if r["situacao"] == "FALTANDO"]
    corrompidos = [r for r in resultados if r["situacao"] == "PDF CORROMPIDO"]
    duplicadas  = [r for r in resultados if r["duplicada"]]

    print("\n" + "=" * 64)
    print("          VERIFICAÇÃO — PLANILHA × PASTA DE NOTAS")
    print("=" * 64)
    print(f"  Linhas na planilha              : {len(resultados)}")
    print(f"  Linhas com nota duplicada       : {len(duplicadas)}")
    print("-" * 64)
    print(f"  ✔  OK (arquivo próprio na pasta): {len(ok)}")
    print(f"  ⚠  FALTANDO                     : {len(faltando)}")
    print(f"  ✘  PDF corrompido               : {len(corrompidos)}")
    print(f"  PDFs na pasta fora da lista     : {len(orfaos)}")
    print("=" * 64)

    if not faltando and not corrompidos:
        print("\n  ✅ TUDO CERTO: cada linha da planilha tem seu próprio PDF na pasta.")
    else:
        print(f"\n  ❌ {len(faltando) + len(corrompidos)} linha(s) precisam ser (re)baixadas "
              f"— use a opção 5 do menu.")

    if faltando:
        print("\n  FALTANDO:")
        for r in faltando:
            dup = " [DUPLICADA]" if r["duplicada"] else ""
            print(f"    Linha {str(r['linha']):>4} | Ocorrência {r['ocorrencia']} | "
                  f"{r['nome_unico']}.pdf{dup}")
    if corrompidos:
        print("\n  PDF CORROMPIDO:")
        for r in corrompidos:
            print(f"    Linha {str(r['linha']):>4} | Ocorrência {r['ocorrencia']} | "
                  f"{r['nome_unico']}.pdf")
    if orfaos:
        print("\n  PDFs na pasta sem linha na planilha:")
        for nome in orfaos:
            print(f"    {nome}")
    print("=" * 64 + "\n")


# =====================================================================
#                              FLUXOS
# =====================================================================
def fluxo_sap(registros=None):
    if registros is None:
        registros = ler_planilha(ARQUIVO_LISTA)
    if not registros:
        return
    input("\nDeixe o SAP aberto e logado. ENTER para começar (Ctrl+C cancela)...")
    dados = baixar_notas_sap(registros)
    imprimir_relatorio_sap(dados)
    caminho = gerar_relatorio_excel(dados)
    if caminho:
        print(f"Relatório salvo em: {caminho}")


def fluxo_nix():
    if not URL_NIX:
        print("NIX_URL não configurada no .env — etapa do portal web não pode rodar.")
        return
    materiais = ler_materiais()
    if not materiais:
        return
    resultados = processar_materiais_nix(materiais)
    atualizar_relatorio(resultados)
    imprimir_resumo_nix(resultados)


def fluxo_verificar():
    """Roda a verificação e devolve a lista de linhas pendentes."""
    registros = ler_planilha(ARQUIVO_LISTA)
    if not registros:
        return []
    resultados = verificar(registros)
    orfaos = listar_pdfs_orfaos(registros)
    imprimir_resumo_verificacao(resultados, orfaos)
    caminho = gerar_planilha_verificacao(resultados, orfaos)
    if caminho:
        print(f"Planilha de verificação salva em: {caminho}")
    return [r for r in resultados if r["situacao"] in ("FALTANDO", "PDF CORROMPIDO")]


def fluxo_completo():
    fluxo_sap()
    fluxo_nix()
    fluxo_verificar()


def fluxo_reprocessar():
    """Verifica, rebaixa só as pendentes (SAP → NIX) e verifica de novo.
    Substitui o antigo PEGARNOTAS_FALTANTES.xlsx + troca manual do ARQUIVO_LISTA."""
    pendentes = fluxo_verificar()
    if not pendentes:
        return
    resp = input(f"Rebaixar {len(pendentes)} nota(s) pendente(s) agora? (s/n): ").strip().lower()
    if resp != "s":
        return
    # Os PDFs corrompidos são apagados para a captura gravar por cima sem conflito
    for r in pendentes:
        caminho = os.path.join(PASTA_DESTINO, f"{r['nome_unico']}.pdf")
        if os.path.exists(caminho) and not pdf_valido(caminho):
            try:
                os.remove(caminho)
            except Exception:
                pass
    fluxo_sap(pendentes)
    fluxo_nix()
    print("\n>>> Conferindo de novo depois do reprocessamento:")
    fluxo_verificar()


def trocar_planilha():
    global ARQUIVO_LISTA
    print(f"\nPlanilha atual: {ARQUIVO_LISTA}")
    novo = input("Novo caminho (pode arrastar o arquivo para cá; ENTER mantém): ").strip().strip('"')
    if not novo:
        return
    if not os.path.exists(novo):
        print("⚠️  Arquivo não encontrado — mantida a planilha atual.")
        return
    ARQUIVO_LISTA = novo
    print("Planilha trocada para esta sessão.")


OPCOES = {
    "1": ("Baixar notas de SERVIÇO no SAP (NIXNFSE)", fluxo_sap),
    "2": ("Baixar notas de MATERIAL no NIX web", fluxo_nix),
    "3": ("Verificar planilha × pasta", fluxo_verificar),
    "4": ("Fluxo completo (1 → 2 → 3)", fluxo_completo),
    "5": ("Rebaixar só o que está faltando", fluxo_reprocessar),
    "6": ("Trocar a planilha de entrada", trocar_planilha),
}


def menu():
    while True:
        print("\n" + "=" * 64)
        print("              DOWNLOAD DE NOTAS — MENU")
        print("=" * 64)
        print(f"  Planilha: {os.path.basename(ARQUIVO_LISTA)}")
        print(f"  Pasta   : {PASTA_DESTINO}")
        print("-" * 64)
        for chave, (descricao, _) in OPCOES.items():
            print(f"  {chave} - {descricao}")
        print("  0 - Sair")
        escolha = input("\nEscolha uma opção: ").strip()

        if escolha == "0":
            print("Até mais!")
            break
        if escolha not in OPCOES:
            print("Opção inválida.")
            continue

        descricao, funcao = OPCOES[escolha]
        print(f"\n>>> {descricao}\n")
        try:
            funcao()
        except KeyboardInterrupt:
            print("\n\n⏹  Interrompido (Ctrl+C) — voltando ao menu.")
        except Exception:
            print("\n✘ Erro inesperado nesta etapa:")
            traceback.print_exc()
            print("Voltando ao menu.")


if __name__ == "__main__":
    menu()